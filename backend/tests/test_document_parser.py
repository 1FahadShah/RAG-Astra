
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pymupdf
from docx import Document as WordDocument

from app.services.document_parser import (
    DocumentParsingError,
    EmptyExtractedText,
    InvalidDocumentForParsing,
    ScannedPdfError,
    UnsupportedDocumentFormat,
    extract_document,
)


class DocumentParserTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.storage_dir = Path(self.temp_dir.name)

        self.settings_patcher = patch(
            "app.services.document_parser.get_settings",
            return_value=SimpleNamespace(storage_dir=self.storage_dir),
        )
        self.settings_patcher.start()
        self.addCleanup(self.settings_patcher.stop)
        self.addCleanup(self.temp_dir.cleanup)

    def test_extracts_utf8_text_file(self):
        path = self.storage_dir / "sample.txt"
        path.write_text(
            "  First line  \r\nSecond line  ",
            encoding="utf-8",
        )

        result = extract_document("sample.txt", "original.txt")

        self.assertEqual(result.source_filename, "original.txt")
        self.assertEqual(result.source_format, "txt")
        self.assertEqual(len(result.units), 1)
        self.assertEqual(
            result.full_text,
            "First line\nSecond line",
        )

    def test_preserves_markdown_syntax(self):
        path = self.storage_dir / "guide.md"
        path.write_text(
            "# Getting Started\n\n- Install Python\n- Run tests",
            encoding="utf-8",
        )

        result = extract_document("guide.md")

        self.assertIn("# Getting Started", result.full_text)
        self.assertIn("- Install Python", result.full_text)
        self.assertEqual(result.source_format, "md")

    def test_extracts_pdf_with_page_provenance(self):
        path = self.storage_dir / "sample.pdf"

        with pymupdf.open() as pdf:
            first = pdf.new_page()
            first.insert_text((72, 72), "First page content")
            second = pdf.new_page()
            second.insert_text((72, 72), "Second page content")
            pdf.save(path)

        result = extract_document("sample.pdf")

        self.assertEqual(len(result.units), 2)
        self.assertEqual(result.units[0].page_number, 1)
        self.assertEqual(result.units[1].page_number, 2)
        self.assertIn("First page content", result.full_text)
        self.assertIn("Second page content", result.full_text)

    def test_pdf_page_numbers_survive_blank_pages(self):
        path = self.storage_dir / "blank-first.pdf"

        with pymupdf.open() as pdf:
            pdf.new_page()
            second = pdf.new_page()
            second.insert_text((72, 72), "Content on page two")
            pdf.save(path)

        result = extract_document("blank-first.pdf")

        self.assertEqual(len(result.units), 1)
        self.assertEqual(result.units[0].page_number, 2)

    def test_rejects_pdf_without_extractable_text(self):
        path = self.storage_dir / "image-only.pdf"

        with pymupdf.open() as pdf:
            pdf.new_page()
            pdf.save(path)

        with self.assertRaises(ScannedPdfError):
            extract_document("image-only.pdf")

    def test_extracts_docx_headings_paragraphs_and_tables(self):
        path = self.storage_dir / "sample.docx"

        document = WordDocument()
        document.add_heading("Introduction", level=1)
        document.add_paragraph("Overview paragraph.")

        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Value"

        document.add_heading("Details", level=2)
        document.add_paragraph("Detailed paragraph.")
        document.save(path)

        result = extract_document("sample.docx")

        self.assertEqual(result.source_format, "docx")
        self.assertIn("Overview paragraph.", result.full_text)
        self.assertIn("Name | Value", result.full_text)
        self.assertIn("Detailed paragraph.", result.full_text)

        details_unit = next(
            unit
            for unit in result.units
            if unit.text == "Detailed paragraph."
        )
        self.assertEqual(details_unit.section_title, "Details")

        indices = [unit.unit_index for unit in result.units]
        self.assertEqual(indices, list(range(len(indices))))

    def test_rejects_empty_text_file(self):
        (self.storage_dir / "empty.txt").write_text(
            "",
            encoding="utf-8",
        )

        with self.assertRaises(EmptyExtractedText):
            extract_document("empty.txt")

    def test_rejects_invalid_utf8(self):
        (self.storage_dir / "invalid.txt").write_bytes(b"\xff\xfe\xfa")

        with self.assertRaises(InvalidDocumentForParsing):
            extract_document("invalid.txt")

    def test_rejects_corrupt_pdf(self):
        (self.storage_dir / "corrupt.pdf").write_bytes(
            b"%PDF-this-is-not-a-valid-pdf"
        )

        with self.assertRaises(InvalidDocumentForParsing):
            extract_document("corrupt.pdf")

    def test_rejects_unsupported_format(self):
        (self.storage_dir / "sample.csv").write_text(
            "a,b\n1,2",
            encoding="utf-8",
        )

        with self.assertRaises(UnsupportedDocumentFormat):
            extract_document("sample.csv")

    def test_rejects_path_traversal(self):
        with self.assertRaises(DocumentParsingError):
            extract_document("../outside.txt")

    def test_rejects_missing_stored_file(self):
        with self.assertRaises(DocumentParsingError):
            extract_document("missing.txt")


if __name__ == "__main__":
    unittest.main()
